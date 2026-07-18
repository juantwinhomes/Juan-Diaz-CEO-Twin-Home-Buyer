#!/usr/bin/env python3
"""Build the red/amber/green QA matrix spreadsheet from a sweep's
qa_results.jsonl (written by reibb_qa_sweep.py).

Usage: python3 reibb_qa_matrix_xlsx.py <qa_results.jsonl> [output.xlsx]

One row per lead, one column per playbook field (playbook section order),
each cell colored: red = required missing/wrong, amber = conditional to
confirm, green = passes. Import into Google Sheets keeps the colors.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else SRC.rsplit(".", 1)[0] + "_matrix.xlsx"

rows = [json.loads(l) for l in open(SRC)]
if not rows:
    print("no leads in results file", file=sys.stderr)
    sys.exit(1)
FIELDS = [m[0] for m in rows[0]["matrix"]]

wb = Workbook()
ws = wb.active
ws.title = "Lead QA Matrix"

hfill = PatternFill("solid", fgColor="1F3864")
hfont = Font(name="Arial", bold=True, color="FFFFFF", size=9)
cfont = Font(name="Arial", size=9)
bold9 = Font(name="Arial", bold=True, size=9)
red = PatternFill("solid", fgColor="F4CCCC")
amber = PatternFill("solid", fgColor="FCE5CD")
green = PatternFill("solid", fgColor="D9EAD3")
thin = Border(*[Side(style="thin", color="BFBFBF")] * 4)
wrap = Alignment(vertical="top", wrap_text=True)

headers = ["Lead (Owner)"] + FIELDS + ["Verdict", "Record Link"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font, cell.fill, cell.border = hfont, hfill, thin
    cell.alignment = Alignment(vertical="center", wrap_text=True)
ws.freeze_panes = "B2"

for r, lead in enumerate(rows, 2):
    name = lead["name"] or "Unknown (phone-only)"
    owner = lead["owner"] or "Unassigned"
    cell = ws.cell(row=r, column=1, value=f"{name}\n({owner})")
    cell.font, cell.border, cell.alignment = bold9, thin, wrap
    for c, (fname, val, status) in enumerate(lead["matrix"], 2):
        cell = ws.cell(row=r, column=c, value=val if val else "MISSING")
        cell.font, cell.border, cell.alignment = cfont, thin, wrap
        cell.fill = {"FAIL": red, "CHECK": amber, "PASS": green}[status]
        if not val:
            cell.font = Font(name="Arial", size=9, bold=True,
                             color="990000" if status == "FAIL" else "7F6000")
    vc = ws.cell(row=r, column=len(FIELDS) + 2, value=lead["verdict"])
    vc.font = Font(name="Arial", bold=True, size=9,
                   color="990000" if lead["verdict"] == "FAIL" else "38761D")
    vc.border, vc.alignment = thin, Alignment(vertical="top")
    lc = ws.cell(row=r, column=len(FIELDS) + 3, value=lead["url"])
    lc.font, lc.border, lc.alignment = cfont, thin, wrap
    ws.row_dimensions[r].height = 64

ws.column_dimensions["A"].width = 20
for i in range(2, len(FIELDS) + 2):
    ws.column_dimensions[get_column_letter(i)].width = 15
ws.column_dimensions[get_column_letter(len(FIELDS) + 2)].width = 9
ws.column_dimensions[get_column_letter(len(FIELDS) + 3)].width = 40
for label, w in {"Tags": 24, "Notes": 26, "Property Address": 24,
                 "Social Profile / URL": 22, "Lead Stage": 18,
                 "Call Disposition": 18}.items():
    if label in FIELDS:
        ws.column_dimensions[get_column_letter(FIELDS.index(label) + 2)].width = w

legend_row = len(rows) + 3
legends = [("Legend:", None),
           ("RED = required field missing or wrong — must fix before closeout", "F4CCCC"),
           ("AMBER = conditional — attach if available / confirm N/A", "FCE5CD"),
           ("GREEN = meets the playbook standard", "D9EAD3")]
for i, (txt, color) in enumerate(legends):
    cell = ws.cell(row=legend_row + i, column=1, value=txt)
    cell.font = Font(name="Arial", size=9, bold=(i == 0), italic=(i > 0))
    if color:
        ws.cell(row=legend_row + i, column=2).fill = PatternFill("solid", fgColor=color)
note = ws.cell(row=legend_row + len(legends) + 1, column=1,
               value="Data written by the REI Lead QA Bot each sweep — no formulas. "
                     "On a Category/Stage/Disposition story mismatch, the fields "
                     "needing change are the red ones.")
note.font = Font(name="Arial", italic=True, size=8)

wb.save(OUT)
print(f"saved {OUT}: {len(rows)} leads x {len(FIELDS)} fields")
