"""Generate semi-monthly payslips from the Equity Track payslip template.

Reads a payroll sheet export (e.g. "Iriga Office June 1 to 15 2026 Payroll - Core Team 1.csv")
and writes one workbook for that cutoff with one payslip tab per employee.

The layout, colors, fonts, merged cells and summary formulas are copied from the
template untouched. Only these are filled in:
  - pay period, employee ID / name / department / position, hourly rate
  - each day's date, hours worked, daily rate, over-time, night diff, gross pay
  - SSS, PhilHealth, Pag-IBIG, withholding tax (and allowance, when the payroll has one)
Monthly Gross Pay Rate keeps the template formula: hourly rate x 8 x 20.

Usage:
    python generate_payslips.py --template Payslip_Template.xlsx \
        --database Employee_Database.csv --payroll "June 1 to 15 Payroll.csv" \
        --out "Payslips June 1-15, 2026.xlsx"
"""
import argparse
import calendar
import csv
import datetime as dt
import re
from copy import copy
from io import BytesIO

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.styles import PatternFill

FILL_REGULAR = "FFFF9900"     # orange (legend: Regular Holiday)
FILL_SPECIAL = "FFF1C232"     # gold (legend: Special Holiday)

# Philippine holidays 2026 (Proclamation No. 1006, s. 2025). Only used for the
# row color + remark; the pay itself comes from the payroll sheet.
HOLIDAYS = {
    dt.date(2026, 1, 1): "Regular Holiday",
    dt.date(2026, 2, 17): "Special Holiday",
    dt.date(2026, 4, 2): "Regular Holiday",
    dt.date(2026, 4, 3): "Regular Holiday",
    dt.date(2026, 4, 4): "Special Holiday",
    dt.date(2026, 4, 9): "Regular Holiday",
    dt.date(2026, 5, 1): "Regular Holiday",
    dt.date(2026, 6, 12): "Regular Holiday",
    dt.date(2026, 8, 21): "Special Holiday",
    dt.date(2026, 8, 31): "Regular Holiday",
    dt.date(2026, 11, 1): "Special Holiday",
    dt.date(2026, 11, 2): "Special Holiday",
    dt.date(2026, 11, 30): "Regular Holiday",
    dt.date(2026, 12, 8): "Special Holiday",
    dt.date(2026, 12, 24): "Special Holiday",
    dt.date(2026, 12, 25): "Regular Holiday",
    dt.date(2026, 12, 30): "Regular Holiday",
    dt.date(2026, 12, 31): "Special Holiday",
}

# Payroll names that differ from the Employee Database spelling
NAME_ALIASES = {
    "Bryan Gene Cadahing": "Bryan Gene Cadahing Hombre",
    "Roiz Emman Bartolata": "Roiz Eman Bartolata",
    "Mc Angelo Abbatuan": "Mc Angelo A. Abbatuan",
    "Kristine Joy Lomeda": "Kristine Joy S. Lomeda",
    "Theavil Margate": "Theavil Marie Margate",
    "Lawrence Oliveros": "John Lawrence Oliveros",
}

# Complete name to print on the payslip when the payroll sheet shortens it
DISPLAY_NAMES = {
    "Bryan Gene Cadahing": "Bryan Gene Cadahing Hombre",
}

# Template geometry ('Dec 16-31' tab = the 16-day version of the template)
SOURCE_SHEET = "Dec 16-31"
SRC_FIRST_DAY_ROW = 12
SRC_DAYS = 16
SRC_TOTAL_ROW = 28
SRC_LAST_ROW = 40
PLAIN_DAY_ROW = 13            # a non-holiday row whose styles are reused for every day
DAY_NAMES = {"Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"}


def money(value):
    s = str(value or "").replace("₱", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:        # blank, "N/A"
        return 0.0


def peso(value):
    return f"₱{value:,.2f}"


def parse_time(text):
    h, m, s = (int(x) for x in text.strip().split(":"))
    return dt.time(h, m, s)


def load_database(path):
    db = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            name = (row.get("Employee Name") or "").strip()
            if name:
                db[name] = {k: (v or "").strip() for k, v in row.items()}
    return db


def next_free_id(used, year):
    nums = [int(m.group(1)) for i in used if (m := re.search(r"(\d{3})-\d{4}$", i))]
    n = max(nums, default=0) + 1
    while True:
        candidate = f"EQT- {n:03d}-{year}"
        if candidate not in used:
            return candidate
        n += 1


def load_payroll(path):
    """Return employee blocks: header totals + one record per calendar day."""
    rows = list(csv.reader(open(path, newline="", encoding="utf-8-sig")))
    department = rows[0][1].strip() if rows and len(rows[0]) > 1 else ""
    employees, cur = [], None
    for r in rows[4:]:
        r = r + [""] * (27 - len(r))
        label = r[2].strip()
        if label and label not in DAY_NAMES:
            cur = {
                "name": label, "rate": r[3].strip(), "hours": r[4].strip(),
                "basic": money(r[5]), "ot": money(r[6]) + money(r[7]), "nd": money(r[8]),
                "gross": money(r[9]), "sss": money(r[13]), "philhealth": money(r[14]),
                "pagibig": money(r[15]), "tax": money(r[17]), "allowance": money(r[18]),
                "net": money(r[19]), "department": department, "days": [],
            }
            employees.append(cur)
        elif label in DAY_NAMES and cur is not None:
            rest_day = money(r[7])
            cur["days"].append({
                "date": dt.datetime.strptime(r[3].strip(), "%m/%d/%Y").date(),
                "hours": parse_time(r[4]),
                "daily": money(r[5]),
                "ot": money(r[6]) + rest_day,
                "nd": money(r[8]),
                "gross": money(r[9]) if r[9].strip() else money(r[5]),
                "rest_day": rest_day,
            })
    return employees


def copy_cell(src, dst):
    if not isinstance(dst, MergedCell):
        dst.value = src.value
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def build_sheet(wb, src, logos, emp, info, days, period_label):
    ws = wb.create_sheet(emp["name"][:31])
    n = len(days)
    shift = n - SRC_DAYS
    last_day = SRC_FIRST_DAY_ROW + n - 1
    total = SRC_TOTAL_ROW + shift

    # Header (rows 1-11) as-is, day rows styled like a plain template row, footer shifted
    for col in range(1, 10):
        for r in range(1, SRC_FIRST_DAY_ROW):
            copy_cell(src.cell(r, col), ws.cell(r, col))
        for i in range(n):
            copy_cell(src.cell(PLAIN_DAY_ROW, col), ws.cell(SRC_FIRST_DAY_ROW + i, col))
        for r in range(SRC_TOTAL_ROW, SRC_LAST_ROW + 1):
            copy_cell(src.cell(r, col), ws.cell(r + shift, col))
    for key, dim in src.column_dimensions.items():
        ws.column_dimensions[key].width = dim.width
    for r in range(1, SRC_LAST_ROW + 1):
        height = src.row_dimensions[r].height
        target = r if r < SRC_FIRST_DAY_ROW else (r + shift if r >= SRC_TOTAL_ROW else None)
        if height and target:
            ws.row_dimensions[target].height = height
    for rng in src.merged_cells.ranges:
        if rng.min_row >= SRC_TOTAL_ROW:
            ws.merge_cells(start_row=rng.min_row + shift, end_row=rng.max_row + shift,
                           start_column=rng.min_col, end_column=rng.max_col)
        else:
            ws.merge_cells(rng.coord)
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines
    for data, width, height, anchor in logos:             # company logo at A1
        logo = Image(BytesIO(data))
        logo.width, logo.height = width, height
        logo.anchor = copy(anchor)
        ws.add_image(logo)
    if src["A1"].hyperlink:
        ws["A1"].hyperlink = src["A1"].hyperlink.target
    ws.sheet_properties = copy(src.sheet_properties)     # fit-to-page
    ws.sheet_format = copy(src.sheet_format)
    ws.page_setup = copy(src.page_setup)
    ws.page_margins = copy(src.page_margins)
    ws.print_options = copy(src.print_options)

    # Employee details
    ws["A4"] = f"Pay Period: {period_label}"
    ws["C6"] = info["id"]
    ws["F6"] = info["department"]
    ws["C7"] = emp["name"]
    ws["F7"] = info["position"]
    ws["C8"] = emp["rate"]
    ws["F8"] = (f'=ROUND(SUMPRODUCT(HOUR(C12:C{last_day}) + MINUTE(C12:C{last_day})/60 + '
                f'SECOND(C12:C{last_day})/3600), 2) & " hours"')

    by_date = {d["date"]: d for d in emp["days"]}
    for i, day in enumerate(days):
        r = SRC_FIRST_DAY_ROW + i
        rec = by_date.get(day, {"hours": dt.time(0, 0), "daily": 0, "ot": 0, "nd": 0,
                                "gross": 0, "rest_day": 0})
        ws[f"A{r}"] = dt.datetime(day.year, day.month, day.day)
        ws[f"B{r}"] = day.strftime("%A")
        ws[f"C{r}"] = rec["hours"]
        ws[f"D{r}"] = peso(rec["daily"])
        ws[f"E{r}"] = peso(rec["ot"])
        ws[f"F{r}"] = peso(rec["nd"])
        ws[f"G{r}"] = peso(rec["gross"])
        holiday = HOLIDAYS.get(day)
        remark = holiday or ("Rest Day OT" if rec["rest_day"] else None)
        ws[f"H{r}"] = remark
        if holiday:
            color = FILL_REGULAR if holiday == "Regular Holiday" else FILL_SPECIAL
            for col in "ABCDEFGH":
                ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=color)

    # Totals + summary (same formulas as the template, pointed at the right rows)
    for col in "DEFG":
        ws[f"{col}{total}"] = (f'="₱" & TEXT(SUMPRODUCT(VALUE(SUBSTITUTE({col}12:{col}{last_day},"₱",""))*1),'
                               f'"#,##0.00")')
    s = total + 4  # Total Basic Pay / Withholding Tax row
    ws[f"D{s}"] = f"=D{total}"
    ws[f"D{s+1}"] = f"=E{total}"
    ws[f"D{s+2}"] = f"=F{total}"
    ws[f"D{s+3}"] = f"=G{total}"
    ws[f"H{s}"] = peso(emp["tax"])
    ws[f"H{s+1}"] = emp["sss"]
    ws[f"H{s+2}"] = emp["philhealth"]
    ws[f"H{s+3}"] = emp["pagibig"]
    for off in (1, 2, 3):
        ws[f"H{s+off}"].number_format = "[$₱]#,##0.00"
    ws[f"H{s+4}"] = f'="₱" & TEXT(SUMPRODUCT(VALUE(SUBSTITUTE(H{s}:H{s+3},"₱",""))*1),"#,##0.00")'
    ws[f"C{s+5}"] = '="₱" & SUBSTITUTE(C8,"₱","")*8*20'
    net = f'VALUE(SUBSTITUTE(D{s+3},"₱","")) - VALUE(SUBSTITUTE(H{s+4},"₱",""))'
    if emp["allowance"]:
        # Same ALLOWANCE row the "August 16 - 31" tab uses, styled like GROSS PAY
        for col in "ABCDE":
            copy_cell(ws[f"{col}{s+3}"], ws[f"{col}{s+4}"])
        for rng in (f"A{s+4}:C{s+4}", f"D{s+4}:E{s+4}"):
            if rng not in ws.merged_cells:
                ws.merge_cells(rng)
        ws[f"A{s+4}"] = "ALLOWANCE:"
        ws[f"D{s+4}"] = emp["allowance"]
        ws[f"D{s+4}"].number_format = "[$₱]#,##0.00"
        net += f" + D{s+4}"
    ws[f"H{s+5}"] = f'="₱" & TEXT({net},"#,##0.00")'
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--database", required=True)
    ap.add_argument("--payroll", required=True)
    ap.add_argument("--out")
    ap.add_argument("--exclude", action="append", default=[],
                    help="payroll name to leave out (repeatable)")
    args = ap.parse_args()

    employees = load_payroll(args.payroll)
    dates = sorted({d["date"] for e in employees for d in e["days"]})
    year, month = dates[0].year, dates[0].month
    first, last = (1, 15) if dates[0].day <= 15 else (16, calendar.monthrange(year, month)[1])
    days = [dt.date(year, month, d) for d in range(first, last + 1)]
    month_name = calendar.month_name[month]
    period_label = f"{month_name} {first} - {last}, {year}"

    db = load_database(args.database)
    used_ids = {row["Employee ID"] for row in db.values()}
    wb = openpyxl.load_workbook(args.template, rich_text=True)
    src = wb[SOURCE_SHEET]
    logos = [(img._data(), img.width, img.height, img.anchor) for img in src._images]
    original = list(wb.worksheets)
    notes = []
    employees = [e for e in employees if e["name"] not in args.exclude]
    for emp in employees:
        emp["payroll_name"] = emp["name"]
        emp["name"] = DISPLAY_NAMES.get(emp["name"], emp["name"])
        row = db.get(NAME_ALIASES.get(emp["payroll_name"], emp["payroll_name"]))
        if row:
            info = {"id": row["Employee ID"], "department": row["Department"], "position": row["Position"]}
        else:
            new_id = next_free_id(used_ids, year)
            used_ids.add(new_id)
            info = {"id": new_id, "department": emp["department"], "position": ""}
            notes.append(f"NOT IN DATABASE: {emp['name']} -> assigned {new_id}, position left blank")
        build_sheet(wb, src, logos, emp, info, days, period_label)

        day_gross = sum(d["gross"] for d in emp["days"])
        if abs(day_gross - emp["gross"]) > 0.01:
            notes.append(f"CHECK {emp['name']}: payroll gross {peso(emp['gross'])} but daily rows add up "
                         f"to {peso(day_gross)} (payslip uses the daily rows)")
    for sheet in original:
        wb.remove(sheet)

    out = args.out or f"Payslips {month_name} {first}-{last}, {year}.xlsx"
    wb.save(out)
    print(f"Saved {out}: {len(employees)} payslips for {period_label}")
    for note in notes:
        print("  " + note)


if __name__ == "__main__":
    main()
