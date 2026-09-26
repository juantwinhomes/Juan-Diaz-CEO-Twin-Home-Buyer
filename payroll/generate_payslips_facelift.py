"""Generate semi-monthly payslips in the August 2026 "facelift" payslip design.

Same inputs and rules as generate_payslips.py (payroll sheet export + employee database);
only the look differs. The style source is any tab of a facelift payslip workbook
(15 daily rows, e.g. "Payslips Iriga August 16-31 2026 Complete.xlsx").

Usage:
    python generate_payslips_facelift.py --design "Payslips ... Complete.xlsx" \
        --database Employee_Database.csv --payroll "August 16 to 31 Payroll.csv" \
        [--logo logo.png] [--out "Payslips August 16-31, 2026.xlsx"]
"""
import argparse
import calendar
import datetime as dt
from copy import copy

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font

from generate_payslips import (DISPLAY_NAMES, HOLIDAYS, NAME_ALIASES, copy_cell, load_database,
                               load_payroll, next_free_id, peso)

# Facelift geometry (15-day source tab)
FIRST_DAY_ROW = 15
SRC_DAYS = 15
SRC_TOTAL_ROW = 30
SRC_LAST_ROW = 38
PLAIN_DAY_ROW = 16
FONT_NAME = "Calibri"          # the design's Carlito is a Calibri look-alike most PCs don't have


def hms(seconds):
    return f"{seconds // 3600}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def build_sheet(wb, src, logo_path, emp, info, days, period_label, extras=False):
    ws = wb.create_sheet("new")          # renamed after the design's own tabs are removed
    ws._final_title = emp["name"][:31]
    n = len(days)
    shift = n - SRC_DAYS
    last_day = FIRST_DAY_ROW + n - 1
    total = SRC_TOTAL_ROW + shift

    for col in range(1, 9):
        for r in range(1, FIRST_DAY_ROW):
            copy_cell(src.cell(r, col), ws.cell(r, col))
        for i in range(n):
            copy_cell(src.cell(PLAIN_DAY_ROW, col), ws.cell(FIRST_DAY_ROW + i, col))
        for r in range(SRC_TOTAL_ROW, SRC_LAST_ROW + 1):
            copy_cell(src.cell(r, col), ws.cell(r + shift, col))
    for key, dim in src.column_dimensions.items():
        ws.column_dimensions[key].width = dim.width
    for r in range(1, SRC_LAST_ROW + 1):
        h = src.row_dimensions[r].height
        if not h:
            continue
        if r < FIRST_DAY_ROW:
            ws.row_dimensions[r].height = h
        elif r >= SRC_TOTAL_ROW:
            ws.row_dimensions[r + shift].height = h
    for i in range(n):
        ws.row_dimensions[FIRST_DAY_ROW + i].height = src.row_dimensions[PLAIN_DAY_ROW].height
    for rng in src.merged_cells.ranges:
        off = shift if rng.min_row >= SRC_TOTAL_ROW else 0
        ws.merge_cells(start_row=rng.min_row + off, end_row=rng.max_row + off,
                       start_column=rng.min_col, end_column=rng.max_col)
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines
    ws.page_margins = copy(src.page_margins)
    if extras:
        ws.page_setup.orientation = "portrait"
        ws.sheet_properties.pageSetUpPr.fitToPage = True  # one page per payslip
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.print_options.horizontalCentered = True

    if extras and logo_path:                               # logo in the empty A1:C5 block
        logo = Image(logo_path)
        logo.height = 125
        logo.width = round(125 * 673 / 537)
        ws.add_image(logo, "A1")

    ws["A7"] = f"Pay Period: {period_label}"
    ws["B9"] = info["id"]
    ws["F9"] = info["department"]
    ws["B10"] = emp["name"]
    ws["F10"] = info["position"] or None
    ws["B11"] = float(emp["rate"].replace("₱", "").replace(",", ""))
    if extras:
        ws["B11"].alignment = Alignment(horizontal="left")

    by_date = {d["date"]: d for d in emp["days"]}
    seconds = 0
    for i, day in enumerate(days):
        r = FIRST_DAY_ROW + i
        rec = by_date.get(day, {"hours": dt.time(0, 0), "daily": 0, "ot": 0, "nd": 0,
                                "gross": 0, "rest_day": 0})
        t = rec["hours"]
        seconds += t.hour * 3600 + t.minute * 60 + t.second
        ws[f"A{r}"] = day.strftime("%m/%d/%Y")
        ws[f"B{r}"] = day.strftime("%A")
        ws[f"C{r}"] = f"{t.hour}:{t.minute:02d}:{t.second:02d}"
        ws[f"D{r}"] = rec["daily"]
        ws[f"E{r}"] = rec["ot"]
        ws[f"F{r}"] = rec["nd"]
        ws[f"G{r}"] = rec["gross"]
        if extras:
            ws[f"H{r}"] = HOLIDAYS.get(day) or ("Rest Day OT" if rec["rest_day"] else None)
        else:
            ws[f"H{r}"] = None
    ws["F11"] = hms(seconds)

    for col in "DEFG":
        ws[f"{col}{total}"] = f"=ROUND(SUM({col}{FIRST_DAY_ROW}:{col}{last_day}),2)"
    s = total + 3                                          # Total Basic Pay / Withholding Tax row
    ws[f"D{s}"] = f"=D{total}"
    ws[f"D{s+1}"] = f"=E{total}"
    ws[f"D{s+2}"] = f"=F{total}"
    ws[f"D{s+3}"] = f"=G{total}"
    ws[f"H{s}"] = emp["tax"]
    ws[f"H{s+1}"] = emp["sss"]
    ws[f"H{s+2}"] = emp["philhealth"]
    ws[f"H{s+3}"] = emp["pagibig"]
    ws[f"D{s+4}"] = emp["allowance"]
    ws[f"H{s+4}"] = f"=ROUND(SUM(H{s}:H{s+3}),2)"
    ws[f"H{s+5}"] = f"=ROUND(D{s+3}+D{s+4}-H{s+4},2)"
    if not extras:
        return ws
    # Monthly Gross Pay Rate (hourly rate x 8 x 20), styled like the allowance line
    for col in "ABCDE":
        copy_cell(ws[f"{col}{s+4}"], ws[f"{col}{s+5}"])
    ws[f"A{s+5}"] = "Monthly Gross Pay Rate:"
    ws[f"D{s+5}"] = "=ROUND(B11*8*20,2)"

    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name == "Carlito":
                f = copy(c.font)
                c.font = Font(name=FONT_NAME, sz=f.sz, b=f.b, i=f.i, color=f.color, u=f.u)
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--design", required=True)
    ap.add_argument("--database", required=True)
    ap.add_argument("--payroll", required=True)
    ap.add_argument("--sheet", default="Core Team")
    ap.add_argument("--logo")
    ap.add_argument("--extras", action="store_true",
                    help="add logo, Monthly Gross Pay Rate, remarks, Calibri font and fit-to-page "
                         "(default: copy the design exactly)")
    ap.add_argument("--out")
    ap.add_argument("--exclude", action="append", default=[])
    args = ap.parse_args()

    employees, orphans = load_payroll(args.payroll, args.sheet)
    notes = [f"SKIPPED unnamed block of {len(o['days'])} days under {o['after']}" for o in orphans]
    notes += [f"SKIPPED {e['name']}: no daily hours" for e in employees if not e["days"]]
    employees = [e for e in employees if e["name"] not in args.exclude and e["days"]]

    dates = sorted({d["date"] for e in employees for d in e["days"]})
    year, month = dates[0].year, dates[0].month
    first, last = (1, 15) if dates[0].day <= 15 else (16, calendar.monthrange(year, month)[1])
    days = [dt.date(year, month, d) for d in range(first, last + 1)]
    stray = [d for d in dates if d not in days]
    if stray:
        raise SystemExit(f"Payroll has dates outside {first}-{last}: {stray[:5]} - fix the payroll sheet")
    period_label = f"{calendar.month_name[month]} {first} - {last}, {year}"

    db = load_database(args.database)
    used_ids = {row["Employee ID"] for row in db.values()}
    wb = openpyxl.load_workbook(args.design)
    src = wb.worksheets[0]
    original = list(wb.worksheets)
    for emp in employees:
        emp["payroll_name"] = emp["name"]
        emp["name"] = DISPLAY_NAMES.get(emp["name"], emp["name"])
        row = db.get(NAME_ALIASES.get(emp["payroll_name"], emp["payroll_name"]))
        if row:
            info = {"id": row["Employee ID"], "department": row["Department"], "position": row["Position"]}
        else:
            new_id = next_free_id(used_ids, year)
            used_ids.add(new_id)
            info = {"id": new_id, "department": emp["department"], "position": emp["job_role"]}
            notes.append(f"NOT IN DATABASE: {emp['name']} -> assigned {new_id}, "
                         f"position '{emp['job_role'] or '(blank)'}' from payroll")
        build_sheet(wb, src, args.logo, emp, info, days, period_label, args.extras)
        day_gross = sum(d["gross"] for d in emp["days"])
        net = day_gross - emp["sss"] - emp["philhealth"] - emp["pagibig"] - emp["tax"] + emp["allowance"]
        if abs(day_gross - emp["gross"]) > 0.01 or abs(net - emp["net"]) > 0.02:
            notes.append(f"CHECK {emp['name']}: payslip gross {peso(day_gross)} net {peso(net)}, "
                         f"payroll gross {peso(emp['gross'])} net {peso(emp['net'])}")
        basic = sum(d["daily"] + d["ot"] + d["nd"] for d in emp["days"])
        if abs(basic - day_gross) > 0.01:
            notes.append(f"CHECK {emp['name']}: daily rate + OT + ND = {peso(basic)} but gross pay "
                         f"column = {peso(day_gross)}")
    for sheet in original:
        wb.remove(sheet)
    for sheet in wb.worksheets:
        sheet.title = sheet._final_title

    out = args.out or f"Payslips {calendar.month_name[month]} {first}-{last}, {year}.xlsx"
    wb.save(out)
    print(f"Saved {out}: {len(employees)} payslips for {period_label}")
    for note in notes:
        print("  " + note)


if __name__ == "__main__":
    main()
