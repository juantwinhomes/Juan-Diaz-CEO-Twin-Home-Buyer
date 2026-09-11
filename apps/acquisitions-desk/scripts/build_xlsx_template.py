#!/usr/bin/env python3
"""Builds the Master Database template (all 12 tabs, exact headers, seeded SETTINGS / USERS / TOOL_INVENTORY)
as an .xlsx that converts 1:1 into a Google Sheet. Reads headers and seeds straight from src/Config.gs so the
template can never drift from the code."""
import re, json, sys, os, datetime, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Evaluate Config.gs with node to get the exact constants
js = "const vm=require('vm');const fs=require('fs');const ctx=vm.createContext({});vm.runInContext(fs.readFileSync(process.argv[1],'utf8'),ctx);" \
     "console.log(JSON.stringify({SHEETS:ctx.SHEETS,HEADERS:ctx.HEADERS,DEFAULT_SETTINGS:ctx.DEFAULT_SETTINGS,SEED_USERS:ctx.SEED_USERS,SEED_TOOLS:ctx.SEED_TOOLS,PERMISSION_LEVEL:ctx.PERMISSION_LEVEL,APP_VERSION:ctx.APP_VERSION}))"
cfg = json.loads(subprocess.check_output(['node', '-e', js, os.path.join(ROOT, 'src', 'Config.gs')]))
now = datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%S.000Z')

wb = Workbook(); wb.remove(wb.active)
head_font = Font(bold=True, color='FFFFFF'); head_fill = PatternFill('solid', fgColor='17223B')
for key in cfg['SHEETS']:
    name = cfg['SHEETS'][key]; headers = cfg['HEADERS'][key]
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c); cell.font = head_font; cell.fill = head_fill; cell.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(c)].width = max(14, min(40, len(headers[c - 1]) + 4))
    ws.freeze_panes = 'A2'
    rows = []
    if key == 'SETTINGS':
        rows = [[k, v, 'SYSTEM_SETUP', now] for k, v in cfg['DEFAULT_SETTINGS'].items()]
    elif key == 'USERS':
        for u in cfg['SEED_USERS']:
            uid = 'USR-' + re.sub(r'[^A-Z0-9]', '', u['name'].split(' ')[0].upper())
            rows.append([uid, u['name'], u['email'], u['team'], u['role'], 'TRUE' if (u['active'] and u['email']) else 'FALSE', cfg['PERMISSION_LEVEL'][u['role']], now, now])
    elif key == 'TOOL_INVENTORY':
        for t in cfg['SEED_TOOLS']:
            rows.append([t['tool_id'], t['name'], t.get('description', ''), t.get('built_by', ''), t.get('operator', ''), '', 'Unconfirmed', '', '', t.get('cadence', 'Not set'), t.get('link', ''), 'Decide', '', 'Not handed over yet', '', '', now, now])
    for r in rows: ws.append(r)
    # plain-text number format on data area so dates/timestamps stay as written
    for row in ws.iter_rows(min_row=2, max_row=max(2, len(rows) + 1), max_col=len(headers)):
        for cell in row: cell.number_format = '@'
out = os.path.join(ROOT, 'docs', 'THB_Acquisitions_Desk_Production_Database_TEMPLATE.xlsx')
wb.save(out); print(out, 'tabs:', len(wb.sheetnames), 'version', cfg['APP_VERSION'])
