import csv
import io
import openpyxl

wb = openpyxl.load_workbook('bateman_leads_dashboard.xlsx')
fb = wb['Facebook Leads']
gg = wb['Google Leads']
fb_rows = [[c.value if c.value is not None else '' for c in row] for row in fb.iter_rows(min_row=2)]
gg_rows = [[c.value if c.value is not None else '' for c in row] for row in gg.iter_rows(min_row=2)]

HEADERS = ["Date", "Name", "Address", "Phone Number", "Email", "Disposition",
           "Notes", "Source", "Campaign", "Tracking Link"]

DISPOSITIONS = [
    "Qualified | Appointment",
    "Qualified | Contract",
    "Unqualified Seller | Other",
    "Unqualified Seller | Poor Location",
    "Unqualified Seller | Retail",
    "Unqualified Seller | Already Listed",
    "Non-Seller | Other",
    "Non-Seller | SPAM",
    "No Contact",
    "Test Lead",
    "Duplicate Lead",
]

RNG = "{{RANGE}}"       # placeholder for Disposition column range
SRC = "{{SRC_RANGE}}"   # placeholder for Source column range
DTE = "{{DATE_RANGE}}"  # placeholder for Date column range

rows = []
rows.append(["Bateman Collective / Twin Home Buyer -- Lead Tracking Dashboard"])
rows.append([])
rows.append(["Overview"])
rows.append(["Total Leads", f"=COUNTA({DTE})"])
rows.append(["Junk Leads (Test Lead + Duplicate Lead)",
             f'=COUNTIF({RNG},"Test Lead")+COUNTIF({RNG},"Duplicate Lead")'])
rows.append(["Actionable Leads (Total - Junk)", "=B4-B5"])
rows.append(["Qualified Leads (Appointment + Contract)",
             f'=COUNTIF({RNG},"Qualified | Appointment")+COUNTIF({RNG},"Qualified | Contract")'])
rows.append(["Qualified Rate (of Actionable Leads)", "=IFERROR(B7/B6,0)"])
rows.append([])
rows.append(["Leads by Source"])
rows.append(["Facebook", f'=COUNTIF({SRC},"Facebook")'])
rows.append(["Google", f'=COUNTIF({SRC},"Google")'])
rows.append(["Print", f'=COUNTIF({SRC},"Print")'])
rows.append([])
rows.append(["Leads by Disposition"])
for d in DISPOSITIONS:
    rows.append([d, f'=COUNTIF({RNG},"{d}")'])
rows.append(["(No Disposition Set)", f"=COUNTBLANK({RNG})"])
rows.append([])
rows.append(["Google Ads Cost Analysis"])
rows.append(["Google Ads Spend (this period)", 25400])
rows.append(["Google Leads Count", "=B12"])
rows.append(["Qualified Leads from Google",
             f'=COUNTIFS({SRC},"Google",{RNG},"Qualified | Appointment")'
             f'+COUNTIFS({SRC},"Google",{RNG},"Qualified | Contract")'])
rows.append(["Cost Per Lead (Google)", "=IFERROR(B30/B31,0)"])
rows.append(["Cost Per Qualified Lead (Google)", "=IFERROR(B30/B32,0)"])
rows.append([])
rows.append(["Note: Google Ads spend of $25,400 is from Google Ads account screenshots "
             "(2026-07-21): $14.1K for Jun 1-30, 2026 + $11.3K for Jul 1-21, 2026 "
             "(values as rounded/displayed in the Google Ads UI), approximating the "
             "Google Leads date range (6/8-7/20/2026). Replace with an exact billing "
             "figure for more precision."])
rows.append([])
rows.append(["All Leads (Facebook + Google combined)"])
header_row_index = len(rows)  # 0-based index where the column header row will live
rows.append(HEADERS)
for r in fb_rows:
    rows.append(r)
for r in gg_rows:
    rows.append(r)

DATA_START = header_row_index + 2  # 1-indexed row number of first data row
DATA_END = DATA_START + len(fb_rows) + len(gg_rows) - 1
RANGE = f"F{DATA_START}:F{DATA_END}"
SRC_RANGE = f"H{DATA_START}:H{DATA_END}"
DATE_RANGE = f"A{DATA_START}:A{DATA_END}"

for r in rows:
    for i, cell in enumerate(r):
        if isinstance(cell, str):
            r[i] = cell.replace(RNG, RANGE).replace(SRC, SRC_RANGE).replace(DTE, DATE_RANGE)

buf = io.StringIO()
w = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)
for r in rows:
    w.writerow(r)

with open('bateman_leads_dashboard.csv', 'w', newline='') as f:
    f.write(buf.getvalue())

print("data rows:", DATA_START, "-", DATA_END, "| total rows:", len(rows))
